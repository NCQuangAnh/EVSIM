"""Convert one of our generated SVG figures into an editable PowerPoint slide (native shapes).

    python scripts/svg_to_pptx.py docs/figures/fw_fig1_overview.svg [out.pptx]

Each SVG element becomes a native shape: rect -> (rounded) rectangle, circle -> oval, line -> connector (with an
arrowhead when the SVG has marker-end), path / polygon -> freeform, text -> text box (Arial). <g data-name="...">
becomes a named PowerPoint group, so a row, a card or an icon moves as one piece (ungroup to edit inside).
<g data-raster="..."> (the road maps: thousands of segments) is rendered to a PNG with headless Edge and inserted
as a picture. 1 SVG px = 0.75 pt. Only the SVG subset that scripts/make_framework_figs.py emits is supported.
Runs with the system Python (needs python-pptx, lxml, Pillow; the project .venv has no lxml).
"""
import math
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from lxml import etree
from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

NS = "{http://www.w3.org/2000/svg}"
PX = 9525                                           # EMU per SVG px (96 dpi)
ZWSP = chr(0x200B)                                  # zero-width space the figure script puts after subscripts
EDGE = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"
FONTS = {(False, False): "arial.ttf", (True, False): "arialbd.ttf", (False, True): "ariali.ttf", (True, True): "arialbi.ttf"}
_font_cache = {}


def font(bold, italic, size):
    key = (bold, italic, round(size * 4))
    if key not in _font_cache:
        _font_cache[key] = ImageFont.truetype(f"C:/Windows/Fonts/{FONTS[(bold, italic)]}", max(1, round(size * 4)))
    return _font_cache[key]


def emu(v):
    return Emu(int(round(v * PX)))


def color(v):
    if not v or v == "none":
        return None
    v = v.lstrip("#")
    if len(v) == 3:
        v = "".join(c * 2 for c in v)
    return RGBColor.from_string(v.upper())


def unstyle(shape):
    """Drop the theme style reference (no theme shadow / outline); fill and line are set explicitly."""
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def set_line(shape, stroke, width, dash=None, arrow=False):
    col = color(stroke)
    if col is None or width <= 0:
        shape.line.fill.background()
        return
    shape.line.color.rgb = col
    shape.line.width = Pt(width * 0.75)
    if dash:
        a = [float(t) for t in re.split(r"[ ,]+", dash.strip()) if t]
        shape.line.dash_style = MSO_LINE.SQUARE_DOT if a and a[0] <= 2.5 else MSO_LINE.DASH
    if arrow:
        ln = shape.line._get_or_add_ln()
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")


def set_fill(shape, fill):
    col = color(fill)
    if col is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = col


# ---------------- transforms ----------------
def parse_transform(t):
    tx = ty = 0.0
    s = 1.0
    if not t:
        return tx, ty, s
    m = re.search(r"translate\(([-\d.]+)[ ,]+([-\d.]+)\)", t)
    if m:
        tx, ty = float(m.group(1)), float(m.group(2))
    m = re.search(r"scale\(([-\d.]+)\)", t)
    if m:
        s = float(m.group(1))
    return tx, ty, s


def compose(outer, inner):
    (ox, oy, os_), (ix, iy, is_) = outer, inner
    return ox + os_ * ix, oy + os_ * iy, os_ * is_


def apply(tf, x, y):
    tx, ty, s = tf
    return tx + s * x, ty + s * y


# ---------------- path parsing (M L H V Q A Z, absolute and relative) ----------------
def path_contours(d):
    toks = re.findall(r"[MLHVQAZmlhvqaz]|-?\d*\.?\d+(?:e-?\d+)?", d)
    i, cmd = 0, None
    cur = (0.0, 0.0)
    start = cur
    contours, pts, closed = [], [], []

    def num():
        nonlocal i
        v = float(toks[i])
        i += 1
        return v

    while i < len(toks):
        if re.match(r"[A-Za-z]", toks[i]):
            cmd = toks[i]
            i += 1
            if cmd in "Zz":
                if pts:
                    contours.append(pts)
                    closed.append(True)
                pts, cur = [], start
                continue
        rel = cmd.islower()
        c = cmd.upper()
        if c == "M":
            x, y = num(), num()
            if rel:
                x, y = cur[0] + x, cur[1] + y
            if pts:
                contours.append(pts)
                closed.append(False)
            pts, cur, start = [(x, y)], (x, y), (x, y)
            cmd = "l" if rel else "L"
        elif c == "L":
            x, y = num(), num()
            cur = (cur[0] + x, cur[1] + y) if rel else (x, y)
            pts.append(cur)
        elif c == "H":
            x = num()
            cur = (cur[0] + x if rel else x, cur[1])
            pts.append(cur)
        elif c == "V":
            y = num()
            cur = (cur[0], cur[1] + y if rel else y)
            pts.append(cur)
        elif c == "Q":
            cx, cy, x, y = num(), num(), num(), num()
            if rel:
                cx, cy, x, y = cur[0] + cx, cur[1] + cy, cur[0] + x, cur[1] + y
            p0 = cur
            for k in range(1, 9):
                t = k / 8
                pts.append(((1 - t) ** 2 * p0[0] + 2 * (1 - t) * t * cx + t * t * x,
                            (1 - t) ** 2 * p0[1] + 2 * (1 - t) * t * cy + t * t * y))
            cur = (x, y)
        elif c == "A":
            rx, _ry, _rot, large, sweep, x, y = (num() for _ in range(7))
            if rel:
                x, y = cur[0] + x, cur[1] + y
            # circular arc: centre from the chord (enough for the lock icon's semicircle)
            mx, my = (cur[0] + x) / 2, (cur[1] + y) / 2
            half = math.hypot(x - cur[0], y - cur[1]) / 2
            h = math.sqrt(max(rx * rx - half * half, 0.0))
            nx, ny = (-(y - cur[1]) / (2 * half), (x - cur[0]) / (2 * half)) if half else (0, 0)
            sgn = 1 if large != sweep else -1
            ccx, ccy = mx + sgn * h * nx, my + sgn * h * ny
            a0 = math.atan2(cur[1] - ccy, cur[0] - ccx)
            a1 = math.atan2(y - ccy, x - ccx)
            da = a1 - a0
            if sweep and da < 0:
                da += 2 * math.pi
            if not sweep and da > 0:
                da -= 2 * math.pi
            for k in range(1, 13):
                a = a0 + da * k / 12
                pts.append((ccx + rx * math.cos(a), ccy + rx * math.sin(a)))
            cur = (x, y)
        else:
            i += 1
    if pts:
        contours.append(pts)
        closed.append(False)
    return list(zip(contours, closed))


# ---------------- converter ----------------
class Converter:
    def __init__(self, svg_path):
        self.svg_path = Path(svg_path)
        self.root = etree.parse(str(svg_path)).getroot()
        vb = [float(v) for v in self.root.get("viewBox").split()]
        self.W, self.H = vb[2], vb[3]
        self.defs = [etree.tostring(e).decode() for e in self.root.iter(NS + "clipPath")]
        self.markers = {}
        self.tmp = Path(tempfile.mkdtemp(prefix="svg2pptx_"))
        self.n_raster = 0

    def run(self, out):
        prs = Presentation()
        prs.slide_width, prs.slide_height = emu(self.W), emu(self.H)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self.walk(self.root, slide.shapes, (0.0, 0.0, 1.0))
        prs.save(out)

    # element dispatch
    def walk(self, el, shapes, tf):
        for ch in el:
            if not isinstance(ch.tag, str):
                continue
            tag = ch.tag.replace(NS, "")
            if tag in ("defs", "clipPath", "marker"):
                continue
            if tag == "g":
                self.group(ch, shapes, tf)
            elif tag == "rect":
                self.rect(ch, shapes, tf)
            elif tag == "circle":
                self.circle(ch, shapes, tf)
            elif tag == "line":
                self.line(ch, shapes, tf)
            elif tag in ("path", "polygon"):
                self.path(ch, shapes, tf, tag)
            elif tag == "text":
                self.text(ch, shapes, tf)

    def group(self, g, shapes, tf):
        tf2 = compose(tf, parse_transform(g.get("transform")))
        if g.get("data-raster"):
            self.raster(g, shapes, tf2)
            return
        name = g.get("data-name")
        if not name:
            self.walk(g, shapes, tf2)
            return
        grp = shapes.add_group_shape()
        grp.name = name
        self.walk(g, grp.shapes, tf2)
        if len(grp.shapes) == 0:
            grp._element.getparent().remove(grp._element)

    def rect(self, e, shapes, tf):
        x, y = apply(tf, float(e.get("x", 0)), float(e.get("y", 0)))
        s = tf[2]
        w, h = float(e.get("width")) * s, float(e.get("height")) * s
        rx = float(e.get("rx", 0)) * s
        shp = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rx > 0 else MSO_SHAPE.RECTANGLE,
                               emu(x), emu(y), emu(w), emu(h))
        unstyle(shp)
        if rx > 0:
            shp.adjustments[0] = min(0.5, rx / min(w, h))
        set_fill(shp, e.get("fill", "#000"))
        set_line(shp, e.get("stroke"), float(e.get("stroke-width", 1)) * s, e.get("stroke-dasharray"))
        if w >= self.W - 1 and h >= self.H - 1:
            shp.name = "Background"
        if e.get("opacity"):
            self.alpha(shp, float(e.get("opacity")))

    def alpha(self, shp, op):
        sf = shp.fill._xPr.find(qn("a:solidFill"))
        if sf is not None and len(sf):
            a = etree.SubElement(sf[0], qn("a:alpha"))
            a.set("val", str(int(op * 100000)))

    def circle(self, e, shapes, tf):
        cx, cy = apply(tf, float(e.get("cx")), float(e.get("cy")))
        r = float(e.get("r")) * tf[2]
        shp = shapes.add_shape(MSO_SHAPE.OVAL, emu(cx - r), emu(cy - r), emu(2 * r), emu(2 * r))
        unstyle(shp)
        set_fill(shp, e.get("fill", "#000"))
        set_line(shp, e.get("stroke"), float(e.get("stroke-width", 1)) * tf[2])

    def line(self, e, shapes, tf):
        x1, y1 = apply(tf, float(e.get("x1")), float(e.get("y1")))
        x2, y2 = apply(tf, float(e.get("x2")), float(e.get("y2")))
        c = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, emu(x1), emu(y1), emu(x2), emu(y2))
        unstyle(c)
        set_line(c, e.get("stroke"), float(e.get("stroke-width", 1)) * tf[2], e.get("stroke-dasharray"),
                 arrow=bool(e.get("marker-end")))
        c.name = "Arrow" if e.get("marker-end") else "Line"

    def path(self, e, shapes, tf, tag):
        if tag == "polygon":
            nums = [float(v) for v in re.split(r"[ ,]+", e.get("points").strip())]
            contours = [(list(zip(nums[::2], nums[1::2])), True)]
        else:
            contours = path_contours(e.get("d"))
        contours = [([apply(tf, *p) for p in pts], cl) for pts, cl in contours if len(pts) >= 2]
        if not contours:
            return
        (p0, _), = contours[:1]
        b = shapes.build_freeform(p0[0][0], p0[0][1], scale=PX)
        for k, (pts, cl) in enumerate(contours):
            if k:
                b.move_to(*pts[0])
            b.add_line_segments(pts[1:], close=cl)
        shp = b.convert_to_shape()
        unstyle(shp)
        set_fill(shp, e.get("fill", "#000"))
        set_line(shp, e.get("stroke"), float(e.get("stroke-width", 1)) * tf[2], e.get("stroke-dasharray"),
                 arrow=bool(e.get("marker-end")))

    def text(self, e, shapes, tf):
        # runs: (text, subscript?) from the text node and its <tspan> children (dy > 0 = subscript)
        runs = [(e.text or "", False)]
        for ts in e:
            dy = float(ts.get("dy", 0) or 0)
            runs.append((ts.text or "", dy > 0))
            runs.append((ts.tail or "", False))
        runs = [(t.replace(ZWSP, ""), sb) for t, sb in runs if t.replace(ZWSP, "")]
        s = "".join(t for t, _ in runs)
        if not s.strip():
            return
        size = float(e.get("font-size", 12)) * tf[2]
        bold = e.get("font-weight") == "bold"
        italic = e.get("font-style") == "italic"
        anchor = e.get("text-anchor", "start")
        x, y = apply(tf, float(e.get("x")), float(e.get("y")))
        w = font(bold, italic, size).getlength(s) / 4 + 3
        h = size * 1.25
        left = x if anchor == "start" else x - w / 2 if anchor == "middle" else x - w
        top = y - size * 0.98
        rot = re.search(r"rotate\(([-\d.]+)", e.get("transform") or "")
        tb = shapes.add_textbox(emu(left), emu(top), emu(w), emu(h))
        if rot:
            tb.rotation = float(rot.group(1))
        tf_ = tb.text_frame
        tf_.margin_left = tf_.margin_right = tf_.margin_top = tf_.margin_bottom = 0
        tf_.word_wrap = False
        tf_.auto_size = MSO_AUTO_SIZE.NONE
        tf_.vertical_anchor = MSO_ANCHOR.TOP
        p = tf_.paragraphs[0]
        p.alignment = {"start": PP_ALIGN.LEFT, "middle": PP_ALIGN.CENTER, "end": PP_ALIGN.RIGHT}[anchor]
        col = color(e.get("fill", "#000"))
        for t, sb in runs:
            r = p.add_run()
            r.text = t
            r.font.name = "Arial"
            r.font.size = Pt(size * 0.75)
            r.font.bold = bold
            r.font.italic = italic
            if sb:
                r.font._rPr.set("baseline", "-25000")
            if col is not None:
                r.font.color.rgb = col
        tb.name = f"Text: {s[:30]}"

    def raster(self, g, shapes, tf):
        """Render the group alone to a PNG (3x) with headless Edge; its first rect gives the frame."""
        frame = g.find(NS + "rect")
        x, y = float(frame.get("x")), float(frame.get("y"))
        w, h = float(frame.get("width")), float(frame.get("height"))
        k = 3
        body = etree.tostring(g).decode()
        svg = (f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="{x} {y} {w} {h}" width="{w * k:.0f}" '
               f'height="{h * k:.0f}"><defs>{"".join(self.defs)}</defs>{body}</svg>')
        self.n_raster += 1
        src = self.tmp / f"r{self.n_raster}.svg"
        png = self.tmp / f"r{self.n_raster}.png"
        src.write_text(svg, encoding="utf-8")
        subprocess.run([EDGE, "--headless=new", "--disable-gpu", "--hide-scrollbars", f"--screenshot={png}",
                        f"--window-size={w * k:.0f},{h * k:.0f}", src.as_uri()],
                       capture_output=True, timeout=120)
        X, Y = apply(tf, x, y)
        pic = shapes.add_picture(str(png), emu(X), emu(Y), emu(w * tf[2]), emu(h * tf[2]))
        pic.name = "Map (picture)"


def main():
    src = Path(sys.argv[1])
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix(".pptx")
    Converter(src).run(str(out))
    print("saved", out)


if __name__ == "__main__":
    main()
